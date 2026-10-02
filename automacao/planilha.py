"""Leitura e escrita da planilha de controle (Excel).

Aqui ficam também os diagnósticos: em vez de um erro genérico, o script diz
POR QUE não conseguiu ler (arquivo aberto, bloqueio do Windows, arquivo
criptografado/protegido, formato errado...).
"""
import shutil
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from automacao import config


class ErroPlanilha(Exception):
    """Problema ao acessar a planilha, com mensagem pronta para o usuário."""


# --------------------------- utilidades de texto ---------------------------
def sem_acento(texto):
    s = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in s if not unicodedata.combining(c))


def norm(texto):
    """'Concluído ' -> 'concluido' (minúsculo, sem acento, sem espaços nas pontas)."""
    if texto is None or str(texto).strip().lower() == "nan":
        return ""
    return sem_acento(str(texto)).strip().lower()


def achar_coluna(df, *palavras, obrigatoria=True):
    """Acha a coluna pelo nome, tolerando acento, maiúscula e ordem das colunas."""
    for palavra in palavras:
        for c in df.columns:
            if norm(palavra) in norm(c):
                return c
    if obrigatoria:
        raise ErroPlanilha(f"Coluna '{palavras[0]}' não encontrada. "
                           f"Colunas existentes: {list(df.columns)}")
    return None


# ------------------------------- diagnóstico -------------------------------
def _assinatura(arquivo):
    with open(arquivo, "rb") as f:
        return f.read(8)


def diagnosticar(arquivo):
    """Relatório passo a passo de por que o arquivo abre (ou não). Para `--diagnostico`."""
    arquivo = Path(arquivo)
    linhas = [f"Arquivo: {arquivo}"]
    if not arquivo.exists():
        linhas.append("X NÃO EXISTE nesse caminho. Confira o nome e a pasta (data/).")
        return linhas
    linhas.append(f"OK existe ({arquivo.stat().st_size} bytes)")

    try:
        assinatura = _assinatura(arquivo)
        linhas.append("OK leitura permitida pelo Windows")
    except PermissionError as e:
        linhas.append(f"X LEITURA NEGADA ({e}). {_explicar_permissao(e)}")
        return linhas

    if assinatura.startswith(b"PK"):
        linhas.append("OK formato .xlsx normal (sem criptografia)")
    elif assinatura.startswith(b"\xd0\xcf\x11\xe0"):
        linhas.append("X ARQUIVO CRIPTOGRAFADO OU PROTEGIDO POR SENHA/RÓTULO "
                      "(formato OLE). O Python não abre. Salve uma cópia SEM "
                      "proteção em data/ (ou gere a base demo com scripts/gerar_base_demo.py).")
        return linhas
    else:
        linhas.append("X não parece um .xlsx válido (assinatura desconhecida).")
        return linhas

    try:
        with open(arquivo, "r+b"):
            linhas.append("OK escrita permitida (arquivo FECHADO no Excel)")
    except PermissionError as e:
        linhas.append(f"! escrita negada: {_explicar_permissao(e)} "
                      "(para só LER/SIMULAR isso não impede)")

    try:
        abas = pd.read_excel(arquivo, sheet_name=None, dtype=str)
        for nome, df in abas.items():
            linhas.append(f"OK aba '{nome}': {df.shape[0]} linhas x {df.shape[1]} colunas")
    except Exception as e:  # noqa: BLE001 - queremos mostrar qualquer causa
        linhas.append(f"X erro ao ler abas: {type(e).__name__}: {e}")
    return linhas


def _explicar_permissao(e):
    codigo = getattr(e, "winerror", None)
    if codigo == 32:
        return "O arquivo está ABERTO em outro programa (Excel?). Feche-o."
    return ("O Windows negou o acesso: pasta/arquivo protegido por política da "
            "empresa, antivírus/DLP ou arquivo somente leitura. Mova uma cópia "
            "para a pasta data/ do projeto.")


def checar_leitura(arquivo):
    arquivo = Path(arquivo)
    if not arquivo.exists():
        raise ErroPlanilha(f"Arquivo não encontrado: {arquivo}\n"
                           "Gere a base demo (python scripts/gerar_base_demo.py) "
                           "ou coloque a planilha na pasta data/.")
    try:
        assinatura = _assinatura(arquivo)
    except PermissionError as e:
        raise ErroPlanilha(f"Leitura negada. {_explicar_permissao(e)}") from e
    if assinatura.startswith(b"\xd0\xcf\x11\xe0"):
        raise ErroPlanilha("A planilha está criptografada/protegida (senha ou rótulo). "
                           "Rode `python run.py --diagnostico` para detalhes.")


def checar_escrita(arquivo):
    try:
        with open(arquivo, "r+b"):
            pass
    except PermissionError as e:
        raise ErroPlanilha(f"Sem permissão para gravar. {_explicar_permissao(e)}") from e


# --------------------------------- leitura ---------------------------------
def carregar(arquivo):
    try:
        base = pd.read_excel(arquivo, sheet_name=config.ABA_BASE, dtype=str)
        trat = pd.read_excel(arquivo, sheet_name=config.ABA_TRAT, dtype=str)
    except (zipfile.BadZipFile, ValueError, KeyError) as e:
        raise ErroPlanilha(f"Não consegui ler as abas '{config.ABA_BASE}' e "
                           f"'{config.ABA_TRAT}': {type(e).__name__}: {e}") from e
    base.columns = [str(c).strip() for c in base.columns]
    trat.columns = [str(c).strip() for c in trat.columns]
    return base, trat


def carregar_motivos(arquivo):
    """Lista [(codigo, texto)] da aba Motivos; None se a aba não existir."""
    try:
        df = pd.read_excel(arquivo, sheet_name=config.ABA_MOTIVOS, dtype=str)
    except (ValueError, KeyError):
        return None
    df.columns = [str(c).strip() for c in df.columns]
    c_cod = achar_coluna(df, "codigo", obrigatoria=False)
    c_mot = achar_coluna(df, "motivo", obrigatoria=False)
    if not c_cod or not c_mot:
        return None
    df = df.dropna(subset=[c_cod, c_mot])
    return [(str(a).strip(), str(b).strip()) for a, b in zip(df[c_cod], df[c_mot])]


# --------------------------------- escrita ---------------------------------
def backup(arquivo):
    arquivo = Path(arquivo)
    destino = arquivo.with_name(f"backup_{datetime.now():%Y%m%d_%H%M%S}_{arquivo.name}")
    shutil.copy(arquivo, destino)
    return destino


def registrar(arquivo, linhas):
    """Acrescenta linhas [serial, data, canal, status, obs] na aba Tratativas."""
    wb = load_workbook(arquivo)
    ws = wb[config.ABA_TRAT]
    ultima = ws.max_row
    while ultima > 1 and ws.cell(ultima, 1).value in (None, ""):
        ultima -= 1                      # ignora linhas vazias formatadas no fim
    for i, linha in enumerate(linhas, start=ultima + 1):
        for j, valor in enumerate(linha, start=1):
            ws.cell(i, j, valor)
    wb.save(arquivo)
