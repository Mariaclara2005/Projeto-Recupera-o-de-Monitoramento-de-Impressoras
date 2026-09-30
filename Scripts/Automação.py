"""
Automação - Recuperação de Monitoramento de Impressoras (Simpress).
"""
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# ========================= CONFIGURAÇÃO =========================
RAIZ = Path(__file__).resolve().parent.parent      # ProjetoV1 - Teste
ARQUIVO = RAIZ / "src" / "Planilha_Teste_Automacao_Impressoras1.xlsx"
ABA_BASE = "Planilha1"
ABA_TRAT = "Tratativas"

MODO = "ENVIAR"      # "SIMULAR" (só mostra) | "RASCUNHO" (salva no Outlook) | "ENVIAR"
EMAIL_TESTE = ""      # se preenchido, TODOS os e-mails vão para este endereço
CONTA_REMETENTE = ""  # opcional: e-mail da conta do Outlook que deve enviar
LIMITE_EMAILS = 3    # trava de segurança: máx. de e-mails por execução

STATUS_ALVO = ["sem monitoramento"]   # comparação sem maiúsculas/minúsculas
MARCADOR = "[AUTO]"   # identifica, na aba Tratativas, o que a automação enviou
# ================================================================


def col(df, palavra):
    """Acha a coluna pelo nome (tolera acento/maiúscula/ordem)."""
    for c in df.columns:
        if palavra in str(c).lower():
            return c
    raise KeyError(f"Coluna com '{palavra}' não encontrada. Colunas: {list(df.columns)}")


def checar_arquivo():
    if not ARQUIVO.exists():
        raise SystemExit(f"Arquivo não encontrado: {ARQUIVO}\n"
                         "Coloque a planilha na pasta 'src' do projeto.")
    try:
        with open(ARQUIVO, "r+b"):
            pass
    except PermissionError:
        raise SystemExit("A planilha está aberta no Excel. Feche-a e rode de novo.")


def carregar():
    base = pd.read_excel(ARQUIVO, sheet_name=ABA_BASE, dtype=str)
    trat = pd.read_excel(ARQUIVO, sheet_name=ABA_TRAT, dtype=str)
    base.columns = [str(c).strip() for c in base.columns]
    trat.columns = [str(c).strip() for c in trat.columns]
    return base, trat


def mapear(base, trat):
    """Retorna as impressoras elegíveis, agrupadas por contato (e-mail)."""
    c_serial, c_status = col(base, "serial"), col(base, "status")
    c_email, c_nome = col(base, "email"), col(base, "nome")
    c_end = col(base, "endere")
    c_serial_t, c_obs_t = col(trat, "serial"), col(trat, "observ")

    # "Uma única vez": pula quem já recebeu um envio automático
    ja_enviados = set(
        trat.loc[trat[c_obs_t].fillna("").str.contains(MARCADOR, regex=False), c_serial_t]
        .str.strip()
    )

    sem_mon = base[base[c_status].fillna("").str.strip().str.lower().isin(STATUS_ALVO)].copy()
    print(f"Impressoras SEM monitoramento na base: {len(sem_mon)}")

    sem_mon["_email"] = sem_mon[c_email].fillna("").str.strip().str.lower()
    invalidos = sem_mon[~sem_mon["_email"].str.contains("@")]
    if len(invalidos):
        print(f"  ! {len(invalidos)} sem e-mail válido (ignoradas): "
              f"{', '.join(invalidos[c_serial])}")
    sem_mon = sem_mon[sem_mon["_email"].str.contains("@")]

    ja = sem_mon[sem_mon[c_serial].str.strip().isin(ja_enviados)]
    if len(ja):
        print(f"  - {len(ja)} já receberam mensagem automática (ignoradas): "
              f"{', '.join(ja[c_serial])}")
    fila = sem_mon[~sem_mon[c_serial].str.strip().isin(ja_enviados)]

    grupos = []
    for email, g in fila.groupby("_email", sort=False):
        grupos.append({
            "email": email,
            "nome": str(g.iloc[0][c_nome]).strip(),
            "impressoras": [(r[c_serial].strip(), str(r[c_end]).strip())
                            for _, r in g.iterrows()],
        })
    return grupos


def montar_mensagem(nome, impressoras):
    nome = nome if nome and nome.lower() != "nan" else ""
    saudacao = f"Olá, {nome}!" if nome else "Olá!"
    if len(impressoras) == 1:
        s, e = impressoras[0]
        intro = (f"Identificamos que uma impressora Simpress (série {s}), "
                 f"localizada em {e}, está sem comunicação com a nossa central.")
    else:
        lista = "\n".join(f"  - {s} | {e}" for s, e in impressoras)
        intro = ("Identificamos que as impressoras Simpress abaixo estão sem "
                 f"comunicação com a nossa central:\n\n{lista}")
    return f"""{saudacao}😊

{intro}

🧐 Você poderia nos informar se :

1- Loja em reforma

2- Impressora trocada

3- Sem rede

4- Perda IP

5- Sem energia

6- Equipamento removido

7- HP 4003dw perdeu configuração

Basta responder este e-mail com o número da opção.

Obrigado!🖨️

"""


def enviar_outlook(destino, assunto, corpo):
    import win32com.client  # só é importado fora do modo SIMULAR
    outlook = win32com.client.Dispatch("Outlook.Application")
    mail = outlook.CreateItem(0)
    mail.To, mail.Subject, mail.Body = destino, assunto, corpo
    if CONTA_REMETENTE:
        for conta in outlook.Session.Accounts:
            if conta.SmtpAddress.lower() == CONTA_REMETENTE.lower():
                mail.SendUsingAccount = conta
                break
        else:
            raise RuntimeError(f"Conta '{CONTA_REMETENTE}' não encontrada no Outlook.")
    if MODO == "ENVIAR":
        mail.Send()
    else:
        mail.Save()


def registrar(seriais):
    """Adiciona 1 linha por impressora na aba Tratativas."""
    wb = load_workbook(ARQUIVO)
    ws = wb[ABA_TRAT]
    prox = ws.max_row + 1
    obs = f"{MARCADOR} Envio 1" + (" (TESTE)" if EMAIL_TESTE else "")
    for s in seriais:
        ws.cell(prox, 1, s)
        ws.cell(prox, 2, datetime.now().strftime("%d/%m/%Y"))
        ws.cell(prox, 3, "Email")
        ws.cell(prox, 4, "Aguardando Retorno")
        ws.cell(prox, 5, obs)
        prox += 1
    wb.save(ARQUIVO)


def main():
    checar_arquivo()
    base, trat = carregar()
    grupos = mapear(base, trat)[:LIMITE_EMAILS]

    print(f"\nModo: {MODO} | E-mails a enviar: {len(grupos)}"
          + (f" | TODOS redirecionados para {EMAIL_TESTE}" if EMAIL_TESTE else "") + "\n")
    if not grupos:
        print("Nada a enviar.")
        return

    if MODO == "ENVIAR":
        backup = ARQUIVO.with_name(f"backup_{datetime.now():%Y%m%d_%H%M%S}_{ARQUIVO.name}")
        shutil.copy(ARQUIVO, backup)
        print(f"Backup criado: {backup.name}\n")

    for g in grupos:
        seriais = [s for s, _ in g["impressoras"]]
        destino = EMAIL_TESTE or g["email"]
        assunto = "Impressora Simpress sem comunicação" if len(seriais) == 1 \
            else f"Impressoras Simpress sem comunicação ({len(seriais)})"
        corpo = montar_mensagem(g["nome"], g["impressoras"])
        print(f"-> {destino} | {g['nome']} | séries: {', '.join(seriais)}")

        if MODO == "SIMULAR":
            print(corpo + "-" * 60)
            continue
        try:
            enviar_outlook(destino, assunto, corpo)
            if MODO == "ENVIAR":
                registrar(seriais)   # grava logo após cada envio, evita duplicar
                print("   enviado e registrado em Tratativas.")
            else:
                print("   salvo em Rascunhos.")
        except Exception as e:
            print(f"   ERRO: {type(e).__name__}: {e}")


if __name__ == "__main__":
    main()